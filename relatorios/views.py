from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib import messages
from django.contrib.auth.models import User, Group
from django.db.models import Q
from django.http import HttpResponse
from django.db import connections
import openpyxl
from xhtml2pdf import pisa
from django.template.loader import get_template
from .models import VendasRelatorio, DuplicidadeNFRelatorio, ContrachequeRelatorio
from .forms import FiltroRelatorioForm, FiltroDuplicidadeNFForm, CadastroUsuarioForm, CadastroModuloForm, FiltroContrachequeForm


# Relatório meta vendedores
@login_required
def dashboard(request):
    return render(request, 'relatorios/dashboard.html')

@login_required 
def relatorio_vendas(request):
    form = FiltroRelatorioForm(request.GET or None)
    resultados = VendasRelatorio.objects.using('viewsOracle').none()

    if request.GET and form.is_valid():
        resultados = VendasRelatorio.objects.using('viewsOracle').all()
        
        vendedor = form.cleaned_data.get('vendedor')
        if vendedor:
            resultados = resultados.filter(
                Q(codigo_vendedor__icontains=vendedor) | 
                Q(nome_vendedor__icontains=vendedor)
            )

        data_de = form.cleaned_data.get('data_de')
        data_ate = form.cleaned_data.get('data_ate')
        if data_de and data_ate:
            resultados = resultados.filter(data_emissao__range=[data_de, data_ate])

        produto = form.cleaned_data.get('produto')
        if produto:
            resultados = resultados.filter(descricao_produto__icontains=produto)

    # Identifica qual botão foi clicado (HTML, EXCEL ou PDF)
    acao = request.GET.get('acao', 'html')

    # ==========================================
    # EXPORTAR PARA EXCEL
    # ==========================================
    if acao == 'excel':
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="Relatorio_OVS.xlsx"'
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Vendas"
        
        # Cabeçalhos
        ws.append(['Cód. Vend.', 'Vendedor', 'Emissão', 'Cód. Prod.', 'Descrição Produto'])
        
        # Linhas
        for item in resultados:
            ws.append([
                item.codigo_vendedor, 
                item.nome_vendedor, 
                item.data_emissao.strftime('%d/%m/%Y') if item.data_emissao else '', 
                item.codigo_produto, 
                item.descricao_produto
            ])
        
        wb.save(response)
        return response

    # ==========================================
    # EXPORTAR PARA PDF
    # ==========================================
    elif acao == 'pdf':
        template = get_template('relatorios/pdf_vendas.html')
        html = template.render({'resultados': resultados})
        
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="Relatorio_OVS.pdf"'
        
        pisa_status = pisa.CreatePDF(html, dest=response)
        if pisa_status.err:
            return HttpResponse('Ocorreu um erro ao gerar o PDF', status=500)
        return response
    
    resultados_tela = resultados[:500]

    # Se a ação for 'html' (ou qualquer outra coisa), renderiza a ecrã normal
    return render(request, 'relatorios/relatorio_vendas.html', {'form': form, 'resultados': resultados_tela})

# Relatório Duplicidade de NF venda assistida
@login_required
def relatorio_duplicidade_nf(request):
    form = FiltroDuplicidadeNFForm(request.GET or None)
    resultados = DuplicidadeNFRelatorio.objects.using('viewsOracle').none()

    if request.GET and form.is_valid():
        resultados = DuplicidadeNFRelatorio.objects.using('viewsOracle').all()

        # Filtro de Filial 
        filial = form.cleaned_data.get('filial')
        if filial:
            resultados = resultados.filter(filial__icontains=filial)

        # Filtro de Mensagem NF
        mensagem = form.cleaned_data.get('mensagem')
        if mensagem:
            resultados = resultados.filter(mensagem_nota__icontains=mensagem)

        # Filtro de Data
        data_de = form.cleaned_data.get('data_de')
        data_ate = form.cleaned_data.get('data_ate')
        if data_de and data_ate:
            resultados = resultados.filter(data_emissao__range=[data_de, data_ate])
        elif data_de:
            resultados = resultados.filter(data_emissao__gte=data_de)
        elif data_ate:
            resultados = resultados.filter(data_emissao__lte=data_ate)

    acao = request.GET.get('acao', 'html')

    # ==========================================
    # EXPORTAR PARA EXCEL
    # ==========================================
    if acao == 'excel':
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="Duplicidade_NF_OVS.xlsx"'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Duplicidade NF"
        
        ws.append(['Filial', 'Pedido', 'Cliente', 'NF', 'Série', 'Emissão', 'Mensagem'])
        for item in resultados:
            ws.append([
                item.filial, item.numero_pedido, item.cliente, 
                item.numero_nf, item.serie_nf, 
                item.data_emissao.strftime('%d/%m/%Y') if item.data_emissao else '',
                item.mensagem_nota
            ])
        wb.save(response)
        return response

    # ==========================================
    # EXPORTAR PARA PDF
    # ==========================================
    elif acao == 'pdf':
        template = get_template('relatorios/pdf_duplicidade_nf.html')
        html = template.render({'resultados': resultados})
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="Duplicidade_NF_OVS.pdf"'
        pisa_status = pisa.CreatePDF(html, dest=response)
        return response
    
    resultados_tela = resultados[:500]

    return render(request, 'relatorios/duplicidade_nf.html', {'form': form, 'resultados': resultados_tela})

# Trava de segurança: Retorna True só se o utilizador for SuperAdmin
def is_admin(user):
    return user.is_superuser

# O user_passes_test expulsa quem não for admin e manda de volta pro início
@login_required
@user_passes_test(is_admin, login_url='/')
def painel_usuarios(request):
    form_usuario = CadastroUsuarioForm()
    form_modulo = CadastroModuloForm()

    if request.method == 'POST':
        # VERIFICA SE O BOTÃO CLICADO FOI O DE USUÁRIO
        if 'btn_salvar_usuario' in request.POST:
            form_usuario = CadastroUsuarioForm(request.POST)
            if form_usuario.is_valid():
                username = form_usuario.cleaned_data['username']
                email = form_usuario.cleaned_data['email']
                password = form_usuario.cleaned_data['password']
                grupo = form_usuario.cleaned_data['grupo']

                if User.objects.filter(username=username).exists():
                    messages.error(request, f'O utilizador "{username}" já existe no sistema!')
                else:
                    novo_usuario = User.objects.create_user(username=username, email=email, password=password)
                    if grupo:
                        novo_usuario.groups.add(grupo)
                    messages.success(request, f'Utilizador "{username}" criado com sucesso!')
                    return redirect('painel_usuarios')

        # VERIFICA SE O BOTÃO CLICADO FOI O DE MÓDULO
        elif 'btn_salvar_modulo' in request.POST:
            form_modulo = CadastroModuloForm(request.POST)
            if form_modulo.is_valid():
                nome_modulo = form_modulo.cleaned_data['nome']
                
                if Group.objects.filter(name=nome_modulo).exists():
                    messages.error(request, f'O Módulo "{nome_modulo}" já existe!')
                else:
                    Group.objects.create(name=nome_modulo)
                    messages.success(request, f'Módulo "{nome_modulo}" criado com sucesso!')
                    return redirect('painel_usuarios')

    # Busca os dados para preencher as tabelas
    usuarios_cadastrados = User.objects.all().order_by('-is_superuser', 'username')
    modulos_cadastrados = Group.objects.all().order_by('name')

    context = {
        'form_usuario': form_usuario,
        'form_modulo': form_modulo,
        'usuarios_cadastrados': usuarios_cadastrados,
        'modulos_cadastrados': modulos_cadastrados
    }
    return render(request, 'relatorios/painel_usuarios.html', context)

# Relatório Contracheque
@login_required
@user_passes_test(is_admin, login_url='/')
def relatorio_contracheque(request):
    form = FiltroContrachequeForm(request.GET or None)
    resultados = ContrachequeRelatorio.objects.using('viewsOracle').none()

    if request.GET and form.is_valid():
        resultados = ContrachequeRelatorio.objects.using('viewsOracle').all()

        filial = form.cleaned_data.get('filial')
        if filial:
            resultados = resultados.filter(filial__icontains=filial)

        matricula = form.cleaned_data.get('matricula')
        if matricula:
            resultados = resultados.filter(matricula__icontains=matricula)

        data_de = form.cleaned_data.get('data_de')
        data_ate = form.cleaned_data.get('data_ate')
        if data_de and data_ate:
            resultados = resultados.filter(data_pag__range=[data_de, data_ate])

    acao = request.GET.get('acao', 'html')

    if acao == 'excel':
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="Contracheque_OVS.xlsx"'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['Filial', 'Matricula', 'Nome', 'Verba', 'Tipo', 'Valor', 'Data Pag.'])
        for item in resultados:
            ws.append([item.filial, item.matricula, item.nome, item.descricao_verba, item.tipo_verba, item.valor, item.data_pag.strftime('%d/%m/%Y')])
        wb.save(response)
        return response
    
    return render(request, 'relatorios/contracheque.html', {'form': form, 'resultados': resultados})

# Relatório vendas filtrando por marca
def relatorio_vendas_marca(request):
    resultados = []

    if request.method == 'POST':
        data_inicio = request.POST.get('data_inicio', '').replace('-', '')
        data_fim = request.POST.get('data_fim', '').replace('-', '')
        marca = request.POST.get('marca', '').strip().upper()
        
        # Pega a ação clicada no botão (html, excel ou pdf)
        acao = request.POST.get('acao', 'html')

        if data_inicio and data_fim and marca:
            query = """
                SELECT * FROM V_VENDAS_MARCA
                WHERE EMISSAO BETWEEN %s AND %s
                AND RTRIM(MARCA) LIKE %s
            """
            marca_like = f"%{marca}%"

            with connections['viewsOracle'].cursor() as cursor:
                cursor.execute(query, [data_inicio, data_fim, marca_like])
                colunas = [col[0] for col in cursor.description]
                resultados = [dict(zip(colunas, row)) for row in cursor.fetchall()]

            # ==========================================
            # EXPORTAR PARA EXCEL 
            # ==========================================
            if acao == 'excel':
                response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                response['Content-Disposition'] = 'attachment; filename="Vendas_Marca_OVS.xlsx"'
                
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Vendas por Marca"
                
                # Cabeçalhos
                ws.append(['Filial', 'Emissão', 'Vendedor', 'Nota Fiscal', 'Série', 'Produto', 'Marca'])
                
                # Linhas (buscando do dicionário gerado pelo cursor)
                for item in resultados:
                    ws.append([
                        item['FILIAL'], 
                        item['DATA_EMISSAO'], 
                        item['VENDEDOR'], 
                        item['NOTA_FISCAL'], 
                        item['SERIE'], 
                        item['PRODUTO'], 
                        item['MARCA']
                    ])
                
                wb.save(response)
                return response

            # ==========================================
            # EXPORTAR PARA PDF 
            # ==========================================
            elif acao == 'pdf':
                template = get_template('relatorios/pdf_vendas_marca.html')
                html = template.render({'resultados': resultados})
                
                response = HttpResponse(content_type='application/pdf')
                # Mantemos o 'inline' aqui para burlar o aviso de segurança do Chrome!
                response['Content-Disposition'] = 'inline; filename="Vendas_Marca_OVS.pdf"'
                
                pisa_status = pisa.CreatePDF(html, dest=response)
                if pisa_status.err:
                    return HttpResponse('Ocorreu um erro ao gerar o PDF', status=500)
                return response

            # ==========================================
            # EXIBIÇÃO NA TELA (HTML)
            # ==========================================
            # Limita a 500 para não travar o navegador
            resultados_tela = resultados[:500]
            return render(request, 'relatorios/vendas_marca.html', {'resultados': resultados_tela})

    # Renderiza a página em branco caso seja o primeiro acesso (GET)
    return render(request, 'relatorios/vendas_marca.html', {'resultados': resultados})